import openpyxl
import sys
import codecs
import dateutil.parser as parser
import dateutil.tz as tz
import pandas
import urllib

import collections

#import urllib.parse

sys.path.insert(0, '/home/goldfarb/envriProvenance/enes_graph_use_case/prov_templates/provtemplates')

import provconv
import prov.model as prov
import prov as provbase

filename="PNP_20180321.xlsx"

wb=openpyxl.load_workbook(filename)

data=wb["IP2-PNP_Vorseite"]


date=data["c6"].value

counter=7
persons=list()
run=True

while run:
        index="c"+str(counter)
        run=True
        try:
                if data[index].value!=None and data[index].value!="":
                        persons.append(data[index].value)
                else:
                        run=False
        except:
                run=False
                pass
        counter+=1

counter=6
trees=dict()
run=True
fnames=[data["n5"].value,data["o5"].value,data["p5"].value,data["q5"].value]
#fnames=[data["m5"].value, data["n5"].value,data["o5"].value,data["p5"].value,data["q5"].value]
def checkNone(data):
        if data==None:
                return ""
        else:
                return str(data).strip("() ")

            
while run:
        index="m"+str(counter)
        run=True
        try:
                if data[index].value!=None and data[index].value!="":
                        idx=checkNone(str(data[index].value))
                        trees[idx]=dict()
                        tree=trees[idx]
                        tree[fnames[0]]=checkNone(data["n"+str(counter)].value)
                        tree[fnames[1]]=checkNone(data["o"+str(counter)].value)
                        tree[fnames[2]]=checkNone(data["p"+str(counter)].value)
                        tree[fnames[3]]=checkNone(data["q"+str(counter)].value)
                        trees[idx]=tree
                else:
                        run=False
        except:
                run=False
                pass
        counter+=1
        
   

out=dict()
#for n in fnames:
#    output[n]=list()   
extracols=["date", "person","tree"]
cols=extracols+fnames
for col in cols:
    out[col]=[]
    for p in persons:
        for t in trees:
            if col not in extracols:
                try:
                    out[col].append(trees[t][col])
                except:
                    pass
            else:
                if col=="date":
                    out[col].append(str(date.isoformat()))
                if col=="person":
                    out[col].append(p)
                if col=="tree":
                    out[col].append(t)
  
#print(repr(out))
data=pandas.DataFrame(out)                    
#print(data)

    #add manual stuff
    
#think about it another way
#create mapping
#colname varname value type



#manual mapping for colname-varname-type
bindmap={ "Anmerkung" : { "varname" : "comment", "type" : "attr", "val" : "literal", "uniqueOnly" : False},
  "Dendrometer" : { "varname" : "dendrometer", "type" : "entity", "val" : "iri",  "uniqueOnly" : False},
  "[cm]" : { "varname" : "readValue", "type" : "attr", "val" : "int", "uniqueOnly" : False},
  "date" : { "varname" : "endDate", "type" : "attr", "val" : "datetime", "uniqueOnly" : False },
  "person" : { "varname" : "readingAgent", "type" : "entity", "val" : "iri", "uniqueOnly" : True },
  "tree" : { "varname" : "tree", "type" : "entity", "val" : "iri", "uniqueOnly" : False }
 }

template="""
@prefix prov: <http://www.w3.org/ns/prov#> .
@prefix tmpl: <http://openprovenance.org/tmpl#> .
@prefix var: <http://openprovenance.org/var#> .
@prefix ex: <http://example.com#> .

"""

outNSpref="ex"
outNS="http://example.com#"

#bind_dicts=[]
bind_dict=dict()
bindfile_dict=dict()

for index, row in data.iterrows():
    rtemplate=template
    for col in data.columns.values:
        if col in bindmap:

            outval=row[col]

            if bindmap[col]["val"]=="iri":
                outval=prov.QualifiedName(prov.Namespace(outNSpref,outNS), urllib.quote(str(outval.encode('utf8', 'replace'))))

	    #print col + " " + bindmap[col]["varname"] + " " + outval
	    ID = "var:"+bindmap[col]["varname"]
	    #ID = prov.QualifiedName(prov.Namespace("var", "http://openprovenance.org/var#"), urllib.quote(bindmap[col]["varname"]))

	    #prepare data for bindings dict
	
	    if ID not in bind_dict:
            	bind_dict[ID]=outval
	    else:
		if not isinstance(bind_dict[ID], list):
			tmp=list()
			tmp.append(bind_dict[ID])
			bind_dict[ID]=tmp
		if not bindmap[col]["uniqueOnly"] or  outval not in bind_dict[ID]:
			bind_dict[ID].append(outval)
            #outval=row[col]

           # if bindmap[col]["val"]=="literal":
           #     outval='"'+outval+'"'
                #outval='"'+outval+'"^^xsd:string'
   #         if bindmap[col]["val"]=="int":
    #            outval=''+outval+'^^xsd:int'
           # if bindmap[col]["val"]=="datetime":
           #     outval='"'+outval+'"'
                #outval='"'+outval+'"^^xsd:datetime'

	    #prepare data for bindings file 

	    if ID not in bindfile_dict:
		bindfile_dict[ID]=dict()
		bindfile_dict[ID]["value"]=list()
		bindfile_dict[ID]["type"]=bindmap[col]["type"]
	    
	    if not bindmap[col]["uniqueOnly"] or outval not in bindfile_dict[ID]["value"]:
		bindfile_dict[ID]["value"].append(outval)


tmpl_NS=prov.Namespace("tmpl", "http://openprovenance.org/tmpl#")

bindDoc=prov.ProvDocument()
bindDoc.add_namespace("tmpl", "http://openprovenance.org/tmpl#") 
bindDoc.add_namespace("var", "http://openprovenance.org/var#") 
#make bindings file, use prov library for this
for ID in bindfile_dict:
	cnt1=0
	attrs=dict()
	for a in bindfile_dict[ID]["value"]:
		#print ID
		#print a
		if bindfile_dict[ID]["type"]=="attr":
			cnt2=0
			if isinstance(a, list):
				for b in a:
					attr=prov.QualifiedName(tmpl_NS, "2dvalue_" + str(cnt1) + "_" + str(cnt2))
					#attr="tmpl:2dvalue_" + str(cnt1) + "_" + str(cnt2)
					cnt2+=1
					attrs[attr]=b
			else:
				attr=prov.QualifiedName(tmpl_NS, "2dvalue_" + str(cnt1) + "_" + str(cnt2))
				#attr="tmpl:2dvalue_" + str(cnt1) + "_" + str(cnt2)
                                attrs[attr]=a
		else:
			attr=prov.QualifiedName(tmpl_NS, "value_" + str(cnt1))
                        attrs[attr]=a
                cnt1+=1
	#print ID			
	#print attrs
	bindDoc.new_record(prov.PROV_ENTITY, prov.Identifier(ID), attributes=attrs)

bindDoc.new_record(prov.PROV_ENTITY, prov.Identifier("var:dendroPlan"), { "tmpl:value_0" : "ex:thePlan1", "tmpl:value_1" : "ex:thePlan2" })
bindDoc.new_record(prov.PROV_ENTITY, prov.Identifier("var:organization"), { "tmpl:value_0" : "ex:theOrganization"})

outfileBind=open("excelProvTemplate_bin.ttl", "w")
#outfileBind.write(bindDoc.serialize(format="provn"))		
#outfileBind.write(bindDoc.serialize(format="provn", rdf_form="xml"))		
outfileBind.write(bindDoc.serialize(format="rdf", rdf_form="ttl"))		
outfileBind_json=open("excelProvTemplate_bin.json", "w")
#outfileBind.write(bindDoc.serialize(format="provn"))		
#outfileBind.write(bindDoc.serialize(format="provn", rdf_form="xml"))		
outfileBind_json.write(bindDoc.serialize(format="json"))		
    


bind_dict["var:dendroPlan"]=[prov.QualifiedName(prov.Namespace(outNSpref,outNS), "thePlan1"), prov.QualifiedName(prov.Namespace(outNSpref,outNS), "thePlan2")]    
bind_dict["var:organization"]=prov.QualifiedName(prov.Namespace(outNSpref,outNS), "theOrganization")    
    
provtemplate=prov.ProvDocument()    
res=provtemplate.deserialize(source="excelProvTemplate.rdf", format="rdf", rdf_format="xml")
print(bind_dict)
#print(res.serialize(format="rdf"))
exp=provconv.instantiate_template(res, bind_dict)

outfile=open("excelProvTemplate_exp.provn", "w")
outfile.write(exp.serialize(format="provn"))
outfile.close()
